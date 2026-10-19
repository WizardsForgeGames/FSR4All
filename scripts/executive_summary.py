"""
FSR4 Executive Summarizer
=========================
Parses NCU profile CSV + benchmark stdout to produce a dark-themed Excel report
with GPU utilization breakdown, per-kernel metrics, and bottleneck analysis.

Adapted from Forge's executive_summary.py for FSR4 inference profiling.

Usage:
    python scripts/executive_summary.py experiments/ncu_fsr4_details.csv --stdout experiments/benchmark_stdout.txt
    python scripts/executive_summary.py experiments/ncu_fsr4_details.csv -o experiments/fsr4_report.xlsx
"""

import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import PieChart, BarChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.text import RichText
    from openpyxl.chart.title import title_maker
    from openpyxl.drawing.text import (
        Paragraph, ParagraphProperties, CharacterProperties,
        RegularTextRun, Font as DrawingFont,
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl required. Install with: pip install openpyxl")
    sys.exit(1)


# ======================================================================
#  DARK THEME PALETTE (same as Forge executive_summary.py)
# ======================================================================

BG_DEEP     = "0B1017"
BG_BASE     = "111A26"
BG_ALT      = "162030"
BG_HEADER   = "1A3352"
BG_TOTAL    = "1C2E48"
BG_CALLOUT  = "0E1820"
BORDER_CLR  = "243040"
TX_PRIMARY  = "D8E2EC"
TX_BRIGHT   = "F0F4F8"
TX_DIM      = "6B7F94"
TX_CYAN     = "00D4FF"
TX_GOLD     = "FFD54F"

STATUS_GOOD   = ("0D2E1E", "4ADE80")
STATUS_WARN   = ("2C2400", "FBBF24")
STATUS_BAD    = ("301212", "F87171")
STATUS_MEM    = ("0E1C30", "60A5FA")
STATUS_COMP   = ("1C0E30", "C084FC")
STATUS_LAUNCH = ("262608", "FDE047")

NEON = [
    "FF6B6B", "4ECDC4", "45B7D1", "FFE66D", "A78BFA",
    "F472B6", "34D399", "FB923C", "60A5FA", "E879F9",
    "22D3EE", "84CC16", "F97316", "818CF8", "2DD4BF",
    "FACC15", "C084FC",
]

BAR_MEM     = "60A5FA"
BAR_COMP    = "C084FC"
BAR_LAUNCH  = "FDE047"
BAR_BALANCED = "4ECDC4"


# ======================================================================
#  PRE-BUILT STYLE OBJECTS
# ======================================================================

FILL_BASE   = PatternFill("solid", fgColor=BG_BASE)
FILL_ALT    = PatternFill("solid", fgColor=BG_ALT)
FILL_HEADER = PatternFill("solid", fgColor=BG_HEADER)
FILL_DEEP   = PatternFill("solid", fgColor=BG_DEEP)
FILL_TOTAL  = PatternFill("solid", fgColor=BG_TOTAL)
FILL_CALL   = PatternFill("solid", fgColor=BG_CALLOUT)

FONT_DATA    = Font(color=TX_PRIMARY, size=10, name="Consolas")
FONT_HEADER  = Font(color=TX_BRIGHT, size=10, bold=True, name="Consolas")
FONT_TITLE   = Font(color=TX_CYAN, size=16, bold=True, name="Consolas")
FONT_SECTION = Font(color=TX_CYAN, size=12, bold=True, name="Consolas")
FONT_LABEL   = Font(color=TX_DIM, size=10, name="Consolas")
FONT_VALUE   = Font(color=TX_BRIGHT, size=11, bold=True, name="Consolas")
FONT_GOLD    = Font(color=TX_GOLD, size=11, bold=True, name="Consolas")
FONT_DIM     = Font(color=TX_DIM, size=10, name="Consolas")

DARK_BORDER = Border(
    left=Side("thin", color=BORDER_CLR),
    right=Side("thin", color=BORDER_CLR),
    top=Side("thin", color=BORDER_CLR),
    bottom=Side("thin", color=BORDER_CLR),
)

ALIGN_C  = Alignment(horizontal="center", vertical="center")
ALIGN_L  = Alignment(horizontal="left", vertical="center")
ALIGN_R  = Alignment(horizontal="right", vertical="center")
ALIGN_CW = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LW = Alignment(horizontal="left", vertical="center", wrap_text=True)

DEC_FMT = '#,##0.00'
PCT_FMT = '0.0'
INT_FMT = '#,##0'


# ======================================================================
#  HELPER FUNCTIONS
# ======================================================================

def safe_float(val, default=0.0):
    if isinstance(val, (int, float)):
        return float(val)
    return default

def status_fill_font(bg_fg_pair):
    bg, fg = bg_fg_pair
    return (PatternFill("solid", fgColor=bg),
            Font(color=fg, size=10, name="Consolas"))

def grade_style(value, good=60, warn=30):
    if value >= good:
        return status_fill_font(STATUS_GOOD)
    elif value >= warn:
        return status_fill_font(STATUS_WARN)
    elif value > 0:
        return status_fill_font(STATUS_BAD)
    return FILL_BASE, FONT_DATA

def bottleneck_style(bn_str):
    if "Memory" in bn_str:
        return status_fill_font(STATUS_MEM)
    elif "Compute" in bn_str:
        return status_fill_font(STATUS_COMP)
    elif "Launch" in bn_str:
        return status_fill_font(STATUS_LAUNCH)
    return FILL_BASE, FONT_DATA

def classify_bottleneck(sm, mem, occ):
    if sm < 3 and mem < 10:
        bn = "Launch overhead"
    elif mem > sm * 2:
        bn = "Memory bound"
    elif sm > mem * 2:
        bn = "Compute bound"
    else:
        bn = "Balanced"
    if 0 < occ < 20:
        bn += " + Low occ."
    return bn

def dc(ws, row, col, value, fmt=None, fill=None, font=None, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = fill or FILL_BASE
    cell.font = font or FONT_DATA
    cell.border = DARK_BORDER
    cell.alignment = align or ALIGN_L
    if fmt:
        cell.number_format = fmt
    return cell

def dc_num(ws, row, col, value, fmt=DEC_FMT, fill=None, font=None):
    return dc(ws, row, col, value, fmt=fmt, fill=fill, font=font, align=ALIGN_R)

def dc_pct(ws, row, col, value, fill=None, font=None):
    return dc(ws, row, col, round(value, 1), fmt=PCT_FMT,
              fill=fill, font=font, align=ALIGN_C)

def dc_status(ws, row, col, value, pair):
    f, fn = pair
    return dc(ws, row, col, round(value, 1), fmt=PCT_FMT, fill=f, font=fn, align=ALIGN_C)

def header_row(ws, row, headers, widths=None):
    for ci, h in enumerate(headers, 1):
        dc(ws, row, ci, h, fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CW)
    if widths:
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w

def section_title(ws, row, text, ncols=13):
    dc(ws, row, 1, text, fill=FILL_DEEP, font=FONT_SECTION)
    for c in range(2, ncols + 1):
        dc(ws, row, c, None, fill=FILL_DEEP)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = 22

def fill_blank(ws, row, ncols, fill=FILL_BASE):
    for c in range(1, ncols + 1):
        dc(ws, row, c, None, fill=fill)
    ws.row_dimensions[row].height = 8

def banner(ws, row, text, ncols=13):
    dc(ws, row, 1, text, fill=FILL_DEEP, font=FONT_TITLE, align=ALIGN_L)
    for c in range(2, ncols + 1):
        dc(ws, row, c, None, fill=FILL_DEEP)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = 36

def callout_box(ws, row, text, ncols, pair):
    fill, font = pair
    big_font = Font(color=font.color, size=12, bold=True, name="Consolas")
    dc(ws, row, 1, text, fill=fill, font=big_font, align=ALIGN_L)
    for c in range(2, ncols + 1):
        dc(ws, row, c, None, fill=fill)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = 28

def color_bar_chart_points(chart, colors):
    if not chart.series:
        return
    series = chart.series[0]
    for i, color in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = color
        series.data_points.append(pt)

def color_pie_chart_points(chart, n):
    if not chart.series:
        return
    series = chart.series[0]
    for i in range(n):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = NEON[i % len(NEON)]
        series.data_points.append(pt)

def _light_cp(size=900, bold=False, color=TX_PRIMARY):
    return CharacterProperties(
        sz=size, b=bold, solidFill=color,
        latin=DrawingFont(typeface="Consolas"),
    )

def _light_txpr(size=900, color=TX_PRIMARY):
    return RichText(
        p=[Paragraph(pPr=ParagraphProperties(defRPr=_light_cp(size, color=color)))]
    )

def _make_title(text, size=1300, color=TX_CYAN):
    cp = _light_cp(size, bold=True, color=color)
    t = title_maker(text)
    for p in t.tx.rich.paragraphs:
        p.pPr = ParagraphProperties(defRPr=cp)
        for r in p.r:
            r.rPr = cp
    return t

def _styled_datalabels(**kwargs):
    dl = DataLabelList(**kwargs)
    dl.txPr = _light_txpr(900)
    return dl

def style_chart_dark(chart, title_str=None):
    try:
        chart.graphical_properties = GraphicalProperties()
        chart.graphical_properties.solidFill = BG_DEEP
    except Exception:
        pass
    if title_str:
        chart.title = _make_title(title_str)
    for attr in ("x_axis", "y_axis"):
        try:
            axis = getattr(chart, attr, None)
            if axis is not None:
                axis.txPr = _light_txpr(900)
        except Exception:
            pass
    try:
        if chart.legend is not None:
            chart.legend.txPr = _light_txpr(800)
    except Exception:
        pass

def prefill_dark(ws, rows=100, cols=43):
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c).fill = FILL_BASE


# ======================================================================
#  NCU CSV PARSER (generic, same as Forge)
# ======================================================================

def parse_ncu_csv(path):
    kernels = {}
    kernel_order = []
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            kid = int(row["ID"])
            kname = row["Kernel Name"]
            short = re.sub(r"\(.*\)", "", kname).strip()
            metric = row["Metric Name"]
            value = row["Metric Value"]
            unit = row.get("Metric Unit", "")
            key = (kid, short)
            if key not in kernels:
                kernels[key] = {"id": kid, "name": short, "full_name": kname,
                                "block_size": row.get("Block Size", ""),
                                "grid_size": row.get("Grid Size", "")}
                kernel_order.append(key)
            try:
                val = float(value.replace(",", ""))
            except (ValueError, AttributeError):
                val = value
            kernels[key][metric] = val
            kernels[key][f"{metric}_unit"] = unit
    return [kernels[k] for k in kernel_order]


def get_duration_us(k):
    dur = k.get("Duration", 0)
    unit = k.get("Duration_unit", "")
    dur = safe_float(dur)
    if "nsecond" in unit or "ns" in unit.lower():
        return dur / 1000
    elif "msecond" in unit:
        return dur * 1000
    return dur


def aggregate_kernels(kernels):
    agg = defaultdict(lambda: {
        "count": 0, "total_us": 0,
        "sm": [], "mem": [], "occ": [], "l1": [], "l2": [],
        "waves": [], "warp_stall": [], "regs": [],
    })
    for k in kernels:
        name = k["name"]
        agg[name]["count"] += 1
        agg[name]["total_us"] += get_duration_us(k)
        agg[name]["sm"].append(safe_float(k.get("Compute (SM) Throughput")))
        agg[name]["mem"].append(safe_float(k.get("Memory Throughput")))
        agg[name]["occ"].append(safe_float(k.get("Achieved Occupancy")))
        agg[name]["l1"].append(safe_float(k.get("L1/TEX Hit Rate")))
        agg[name]["l2"].append(safe_float(k.get("L2 Hit Rate")))
        agg[name]["waves"].append(safe_float(k.get("Waves Per SM")))
        agg[name]["warp_stall"].append(safe_float(k.get("Warp Cycles Per Executed Instruction")))
        agg[name]["regs"].append(safe_float(k.get("Registers Per Thread")))
    result = []
    for name, v in agg.items():
        n = v["count"]
        avg = lambda lst: sum(lst) / len(lst) if lst else 0
        result.append({
            "name": name, "count": n,
            "total_us": v["total_us"],
            "avg_us": v["total_us"] / n,
            "sm": avg(v["sm"]), "mem": avg(v["mem"]),
            "occ": avg(v["occ"]), "l1": avg(v["l1"]),
            "l2": avg(v["l2"]), "waves": avg(v["waves"]),
            "warp_stall": avg(v["warp_stall"]),
            "regs": avg(v["regs"]),
        })
    result.sort(key=lambda x: -x["total_us"])
    return result


# ======================================================================
#  FSR4 BENCHMARK STDOUT PARSER
# ======================================================================

def parse_fsr4_stdout(path):
    """Parse FSR4 benchmark stdout for per-layer timings and metadata."""
    data = {}
    text = Path(path).read_text(encoding="utf-8", errors="replace")

    # GPU name
    m = re.search(r"GPU:\s+(.+)", text)
    if m: data["gpu"] = m.group(1).strip()

    # Resolution
    m = re.search(r"Resolution:\s+(\d+)x(\d+)", text)
    if m:
        data["width"] = int(m.group(1))
        data["height"] = int(m.group(2))
        data["resolution"] = f"{m.group(1)}x{m.group(2)}"

    # Throughput
    m = re.search(r"Throughput:\s+([\d.]+)\s*FPS", text)
    if m: data["fps"] = float(m.group(1))

    m = re.search(r"Average:\s+([\d.]+)\s*ms/frame", text)
    if m: data["avg_ms"] = float(m.group(1))

    # Memory
    m = re.search(r"Weights:\s+([\d.]+)\s*KB", text)
    if m: data["weights_kb"] = float(m.group(1))

    m = re.search(r"Workspace:\s+([\d.]+)\s*MB", text)
    if m: data["workspace_mb"] = float(m.group(1))

    m = re.search(r"Total VRAM:\s+([\d.]+)\s*MB", text)
    if m: data["total_vram_mb"] = float(m.group(1))

    # Per-layer timings from "Averaged Per-Layer" section
    data["phases"] = {}
    # Find the averaged section
    avg_section = re.search(
        r"Averaged Per-Layer.*?\n(.*?)(?:Memory Usage|$)",
        text, re.DOTALL
    )
    if avg_section:
        for m in re.finditer(r"^\s+(.+?)\s{2,}([\d.]+)\s+[\d.]+%", avg_section.group(1), re.MULTILINE):
            layer_name = m.group(1).strip()
            time_ms = float(m.group(2))
            data["phases"][layer_name] = time_ms

    return data


# ======================================================================
#  EXCEL BUILDER
# ======================================================================

def build_excel(kernels, stdout_data, output_path):
    wb = Workbook()
    agg = aggregate_kernels(kernels) if kernels else []
    total_us = sum(k["total_us"] for k in agg)
    W = 13

    # ==================================================================
    #  Sheet 1: Executive Summary
    # ==================================================================
    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_properties.tabColor = "00D4FF"
    prefill_dark(ws)

    row = 1
    banner(ws, row, "FSR4 INFERENCE PERFORMANCE REPORT", W)
    row += 1
    fill_blank(ws, row, W, FILL_DEEP)
    row += 1

    # -- Model Info --
    section_title(ws, row, "MODEL CONFIGURATION", W)
    row += 1
    info_items = [
        ("Architecture", "U-Net INT8 (7->16->32->64->32->16->8)"),
        ("GPU", stdout_data.get("gpu", "Unknown")),
        ("Resolution", stdout_data.get("resolution", "Unknown")),
        ("Weights", f"{stdout_data.get('weights_kb', 0):.1f} KB"),
        ("Workspace", f"{stdout_data.get('workspace_mb', 0):.1f} MB"),
        ("Total VRAM", f"{stdout_data.get('total_vram_mb', 0):.1f} MB"),
    ]
    for i, (label, val) in enumerate(info_items):
        bg = FILL_ALT if i % 2 else FILL_BASE
        dc(ws, row, 1, label, font=FONT_LABEL, fill=bg)
        dc(ws, row, 2, val, font=FONT_VALUE, fill=bg, align=ALIGN_R)
        for c in range(3, W + 1):
            dc(ws, row, c, None, fill=bg)
        row += 1
    fill_blank(ws, row, W)
    row += 1

    # -- Throughput --
    if "avg_ms" in stdout_data:
        section_title(ws, row, "THROUGHPUT", W)
        row += 1
        perf_items = [
            ("Average frame time", f"{stdout_data['avg_ms']:.2f} ms"),
            ("Throughput", f"{stdout_data.get('fps', 0):.1f} FPS"),
        ]
        for i, (label, val) in enumerate(perf_items):
            bg = FILL_ALT if i % 2 else FILL_BASE
            dc(ws, row, 1, label, font=FONT_LABEL, fill=bg)
            dc(ws, row, 2, val, font=FONT_VALUE, fill=bg, align=ALIGN_R)
            for c in range(3, W + 1):
                dc(ws, row, c, None, fill=bg)
            row += 1
        fill_blank(ws, row, W)
        row += 1

    # -- Per-Layer Breakdown --
    if stdout_data.get("phases"):
        section_title(ws, row, "PER-LAYER BREAKDOWN (CUDA Events)", W)
        row += 1
        phase_headers = ["Layer", "Time (ms)", "% Total"]
        header_row(ws, row, phase_headers + [""] * (W - len(phase_headers)),
                   [42, 14, 12] + [10] * (W - 3))
        row += 1
        phase_data_start = row

        total_phase = sum(stdout_data["phases"].values())
        sorted_phases = sorted(stdout_data["phases"].items(), key=lambda x: -x[1])
        for i, (phase, ms) in enumerate(sorted_phases):
            pct = ms / total_phase * 100 if total_phase > 0 else 0
            bg = FILL_ALT if i % 2 else FILL_BASE
            dc(ws, row, 1, phase, fill=bg, font=FONT_DATA)
            dc_num(ws, row, 2, round(ms, 3), fill=bg)
            if pct > 30:
                dc_status(ws, row, 3, pct, status_fill_font(STATUS_BAD))
            elif pct > 10:
                dc_status(ws, row, 3, pct, status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws, row, 3, pct, fill=bg)
            for c in range(4, W + 1):
                dc(ws, row, c, None, fill=bg)
            row += 1

        # Totals
        dc(ws, row, 1, "TOTAL", fill=FILL_TOTAL, font=FONT_GOLD)
        dc_num(ws, row, 2, round(total_phase, 3), fill=FILL_TOTAL, font=FONT_GOLD)
        dc_pct(ws, row, 3, 100.0, fill=FILL_TOTAL, font=FONT_GOLD)
        for c in range(4, W + 1):
            dc(ws, row, c, None, fill=FILL_TOTAL)
        row += 1

        # Phase pie chart
        chart = PieChart()
        chart.style = 26
        n = len(sorted_phases)
        data_ref = Reference(ws, min_col=2, min_row=phase_data_start - 1,
                             max_row=phase_data_start + n - 1)
        cats_ref = Reference(ws, min_col=1, min_row=phase_data_start,
                             max_row=phase_data_start + n - 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 22
        chart.height = 14
        chart.dataLabels = _styled_datalabels(
            showPercent=True, showCatName=False, showVal=False, showSerName=False)
        if chart.legend:
            chart.legend.position = "b"
        color_pie_chart_points(chart, n)
        style_chart_dark(chart, "Layer Time Distribution")
        ws.add_chart(chart, f"{get_column_letter(6)}{phase_data_start - 1}")
        fill_blank(ws, row, W)
        row += 1

    # -- NCU Summary Stats --
    if agg:
        section_title(ws, row, "NCU KERNEL PROFILING", W)
        row += 1
        dc(ws, row, 1, "Total kernel time", font=FONT_LABEL, fill=FILL_ALT)
        dc(ws, row, 2, f"{total_us:,.1f} us", font=FONT_VALUE, fill=FILL_ALT, align=ALIGN_R)
        dc(ws, row, 3, f"({len(kernels)} launches)", font=FONT_DIM, fill=FILL_ALT)
        for c in range(4, W + 1):
            dc(ws, row, c, None, fill=FILL_ALT)
        row += 1
        dc(ws, row, 1, "Unique kernels", font=FONT_LABEL)
        dc(ws, row, 2, str(len(agg)), font=FONT_VALUE, align=ALIGN_R)
        for c in range(3, W + 1):
            dc(ws, row, c, None)
        row += 1
        fill_blank(ws, row, W)
        row += 1

    # -- Top Hotspots --
    if agg:
        section_title(ws, row, "TOP KERNEL HOTSPOTS", W)
        row += 1
        hs_headers = ["Kernel", "Time (us)", "% Total", "Calls", "Avg (us)",
                       "SM %", "Mem %", "Occ %", "Warp Stall", "Bottleneck"]
        hs_widths = [34, 14, 10, 8, 12, 9, 9, 9, 12, 28]
        while len(hs_headers) < W:
            hs_headers.append("")
            hs_widths.append(8)
        header_row(ws, row, hs_headers, hs_widths)
        hotspot_hdr = row
        row += 1

        top_n = min(len(agg), 10)
        for i, k in enumerate(agg[:top_n]):
            pct = k["total_us"] / total_us * 100 if total_us > 0 else 0
            bn = classify_bottleneck(k["sm"], k["mem"], k["occ"])
            bg = FILL_ALT if i % 2 else FILL_BASE

            dc(ws, row, 1, k["name"], fill=bg, font=FONT_DATA)
            dc_num(ws, row, 2, round(k["total_us"], 2), fill=bg)
            if pct > 30:
                dc_status(ws, row, 3, pct, status_fill_font(STATUS_BAD))
            elif pct > 10:
                dc_status(ws, row, 3, pct, status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws, row, 3, pct, fill=bg)
            dc(ws, row, 4, k["count"], fill=bg, font=FONT_DATA, align=ALIGN_C)
            dc_num(ws, row, 5, round(k["avg_us"], 2), fill=bg)
            for ci, val in [(6, k["sm"]), (7, k["mem"]), (8, k["occ"])]:
                gf, gfn = grade_style(val)
                dc_status(ws, row, ci, val, (gf, gfn))
            if k["warp_stall"] > 100:
                dc_status(ws, row, 9, k["warp_stall"], status_fill_font(STATUS_BAD))
            elif k["warp_stall"] > 30:
                dc_status(ws, row, 9, k["warp_stall"], status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws, row, 9, k["warp_stall"], fill=bg)
            bn_fill, bn_font = bottleneck_style(bn)
            dc(ws, row, 10, bn, fill=bn_fill, font=bn_font)
            for c in range(11, W + 1):
                dc(ws, row, c, None, fill=bg)
            row += 1

        # Kernel pie chart
        chart = PieChart()
        chart.style = 26
        data_ref = Reference(ws, min_col=2, min_row=hotspot_hdr,
                             max_row=hotspot_hdr + top_n)
        cats_ref = Reference(ws, min_col=1, min_row=hotspot_hdr + 1,
                             max_row=hotspot_hdr + top_n)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 24
        chart.height = 16
        chart.dataLabels = _styled_datalabels(
            showPercent=True, showCatName=False, showVal=False,
            showSerName=False, showLeaderLines=True)
        if chart.legend:
            chart.legend.position = "b"
        color_pie_chart_points(chart, top_n)
        style_chart_dark(chart, "Kernel Time Distribution")
        ws.add_chart(chart, f"{get_column_letter(W + 2)}{hotspot_hdr}")
        fill_blank(ws, row, W)
        row += 1

    # -- Assessment --
    row += 1
    section_title(ws, row, "ASSESSMENT", W)
    row += 1

    if stdout_data.get("phases"):
        total_phase = sum(stdout_data["phases"].values())
        # Find ConvNext time
        convnext_ms = sum(v for k, v in stdout_data["phases"].items()
                          if "ConvNext" in k)
        convnext_pct = convnext_ms / total_phase * 100 if total_phase > 0 else 0
        if convnext_pct > 50:
            callout_box(ws, row,
                        f"  BOTTLENECK  ConvNext 16ch blocks = {convnext_pct:.0f}% of total  |  "
                        f"im2col + GEMM workspace overhead dominant", W,
                        status_fill_font(STATUS_BAD))
        elif convnext_pct > 30:
            callout_box(ws, row,
                        f"  WARNING  ConvNext blocks = {convnext_pct:.0f}%  |  "
                        f"Consider direct spatial conv kernel", W,
                        status_fill_font(STATUS_WARN))
        else:
            callout_box(ws, row,
                        f"  GOOD  No single block type dominates", W,
                        status_fill_font(STATUS_GOOD))
        row += 1

    if agg and total_us > 0:
        top = agg[0]
        top_pct = top["total_us"] / total_us * 100
        bn = classify_bottleneck(top["sm"], top["mem"], top["occ"])
        if top_pct > 30:
            callout_box(ws, row,
                        f"  HOTSPOT  {top['name']}  =  {top_pct:.0f}%  |  "
                        f"{bn}  |  Occ: {top['occ']:.0f}%", W,
                        status_fill_font(STATUS_BAD))
            row += 1

    for r in range(row, row + 10):
        fill_blank(ws, r, W)
    ws.freeze_panes = "A2"

    # ==================================================================
    #  Sheet 2: Kernel Breakdown
    # ==================================================================
    if agg:
        ws2 = wb.create_sheet("Kernel Breakdown")
        ws2.sheet_properties.tabColor = "FF6B6B"
        prefill_dark(ws2)

        row = 1
        banner(ws2, row, "KERNEL TIME BREAKDOWN", W)
        row += 1
        fill_blank(ws2, row, W, FILL_DEEP)
        row += 1

        kb_headers = ["Kernel", "Calls", "Total (us)", "Avg (us)", "% Total",
                      "SM %", "Mem %", "Occ %", "L1 Hit %", "L2 Hit %",
                      "Warp Stall", "Regs", "Bottleneck"]
        kb_widths = [34, 7, 14, 12, 10, 9, 9, 9, 10, 10, 12, 7, 28]
        header_row(ws2, row, kb_headers, kb_widths)
        data_start = row
        row += 1

        bar_colors = []
        for i, k in enumerate(agg):
            pct = k["total_us"] / total_us * 100 if total_us > 0 else 0
            bn = classify_bottleneck(k["sm"], k["mem"], k["occ"])
            bg = FILL_ALT if i % 2 else FILL_BASE

            dc(ws2, row, 1, k["name"], fill=bg)
            dc(ws2, row, 2, k["count"], fill=bg, align=ALIGN_C)
            dc_num(ws2, row, 3, round(k["total_us"], 2), fill=bg)
            dc_num(ws2, row, 4, round(k["avg_us"], 2), fill=bg)
            if pct > 30:
                dc_status(ws2, row, 5, pct, status_fill_font(STATUS_BAD))
            elif pct > 10:
                dc_status(ws2, row, 5, pct, status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws2, row, 5, pct, fill=bg)
            for ci, val in [(6, k["sm"]), (7, k["mem"]), (8, k["occ"]),
                            (9, k["l1"]), (10, k["l2"])]:
                gf, gfn = grade_style(val)
                dc_status(ws2, row, ci, val, (gf, gfn))
            if k["warp_stall"] > 100:
                dc_status(ws2, row, 11, k["warp_stall"], status_fill_font(STATUS_BAD))
            elif k["warp_stall"] > 30:
                dc_status(ws2, row, 11, k["warp_stall"], status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws2, row, 11, k["warp_stall"], fill=bg)
            dc(ws2, row, 12, int(k["regs"]) if k["regs"] else "", fill=bg, align=ALIGN_C)
            bn_fill, bn_font = bottleneck_style(bn)
            dc(ws2, row, 13, bn, fill=bn_fill, font=bn_font)

            if "Memory" in bn: bar_colors.append(BAR_MEM)
            elif "Compute" in bn: bar_colors.append(BAR_COMP)
            elif "Launch" in bn: bar_colors.append(BAR_LAUNCH)
            else: bar_colors.append(BAR_BALANCED)
            row += 1

        fill_blank(ws2, row, W)
        row += 1
        dc(ws2, row, 1, "TOTAL", fill=FILL_TOTAL, font=FONT_GOLD)
        dc(ws2, row, 2, sum(k["count"] for k in agg), fill=FILL_TOTAL,
           font=FONT_GOLD, align=ALIGN_C)
        dc_num(ws2, row, 3, round(total_us, 2), fill=FILL_TOTAL, font=FONT_GOLD)
        for c in range(4, W + 1):
            dc(ws2, row, c, None, fill=FILL_TOTAL)
        row += 2

        chart = BarChart()
        chart.type = "bar"
        chart.style = 26
        chart.y_axis.title = None
        chart.x_axis.title = None
        chart.y_axis.scaling.orientation = "maxMin"
        chart.legend = None
        chart.width = 32
        chart.height = max(12, len(agg) * 1.2 + 4)

        n = len(agg)
        data_ref = Reference(ws2, min_col=3, min_row=data_start,
                             max_row=data_start + n)
        cats_ref = Reference(ws2, min_col=1, min_row=data_start + 1,
                             max_row=data_start + n)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        color_bar_chart_points(chart, bar_colors)
        style_chart_dark(chart, "Kernel Time Distribution (us)")
        ws2.add_chart(chart, f"A{row}")
        ws2.freeze_panes = f"A{data_start + 1}"

    # ==================================================================
    #  Sheet 3: Bottleneck Analysis
    # ==================================================================
    if agg:
        ws3 = wb.create_sheet("Bottleneck Analysis")
        ws3.sheet_properties.tabColor = "FBBF24"
        BN_COLS = 11
        prefill_dark(ws3)

        row = 1
        banner(ws3, row, "BOTTLENECK ANALYSIS", BN_COLS)
        row += 1
        fill_blank(ws3, row, BN_COLS, FILL_DEEP)
        row += 1

        bn_headers = ["Kernel", "% Time", "SM %", "Mem %",
                      "Occ %", "L1 Hit %", "L2 Hit %", "Warp Stall",
                      "Waves/SM", "Bottleneck", "Optimization Hint"]
        bn_widths = [34, 10, 9, 9, 9, 10, 10, 12, 10, 28, 42]
        header_row(ws3, row, bn_headers, bn_widths)
        bn_hdr_row = row
        row += 1

        for i, k in enumerate(agg):
            pct = k["total_us"] / total_us * 100 if total_us > 0 else 0
            bn = classify_bottleneck(k["sm"], k["mem"], k["occ"])
            bg = FILL_ALT if i % 2 else FILL_BASE

            dc(ws3, row, 1, k["name"], fill=bg)
            if pct > 30:
                dc_status(ws3, row, 2, pct, status_fill_font(STATUS_BAD))
            elif pct > 10:
                dc_status(ws3, row, 2, pct, status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws3, row, 2, pct, fill=bg)
            for ci, val in [(3, k["sm"]), (4, k["mem"]), (5, k["occ"]),
                            (6, k["l1"]), (7, k["l2"])]:
                gf, gfn = grade_style(val)
                dc_status(ws3, row, ci, val, (gf, gfn))
            if k["warp_stall"] > 100:
                dc_status(ws3, row, 8, k["warp_stall"], status_fill_font(STATUS_BAD))
            elif k["warp_stall"] > 30:
                dc_status(ws3, row, 8, k["warp_stall"], status_fill_font(STATUS_WARN))
            else:
                dc_pct(ws3, row, 8, k["warp_stall"], fill=bg)
            dc(ws3, row, 9, round(k["waves"], 3) if k["waves"] else 0,
               fmt='0.000', fill=bg, align=ALIGN_C)
            bn_fill, bn_font = bottleneck_style(bn)
            dc(ws3, row, 10, bn, fill=bn_fill, font=bn_font)

            # FSR4-specific hints
            hints = []
            if "im2col" in k["name"].lower():
                hints.append("Replace with direct spatial conv")
            if k["waves"] < 0.1 and k["waves"] > 0:
                hints.append("Grid too small")
            if k["warp_stall"] > 100:
                hints.append("Latency stall")
            if "Memory" in bn and k["l1"] < 30 and k["l1"] > 0:
                hints.append("Improve coalescing")
            if k["occ"] < 20 and k["occ"] > 0:
                hints.append("Low occupancy")
                if k["regs"] > 64:
                    hints.append("Reduce regs")
            if "transpose" in k["name"].lower():
                hints.append("Fuse with producer")
            if "quantize" in k["name"].lower():
                hints.append("Fuse with im2col")
            hints = hints[:2]
            hint_str = " | ".join(hints) if hints else "OK"
            dc(ws3, row, 11, hint_str, fill=bg, font=FONT_DATA, align=ALIGN_L)
            row += 1

        row += 2
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "stacked"
        chart.style = 26
        chart.y_axis.scaling.orientation = "maxMin"
        chart.legend.position = "b"
        chart.width = 28
        chart.height = max(12, len(agg) * 1.2 + 4)

        n = len(agg)
        sm_ref = Reference(ws3, min_col=3, min_row=bn_hdr_row,
                           max_row=bn_hdr_row + n)
        mem_ref = Reference(ws3, min_col=4, min_row=bn_hdr_row,
                            max_row=bn_hdr_row + n)
        cats_ref = Reference(ws3, min_col=1, min_row=bn_hdr_row + 1,
                             max_row=bn_hdr_row + n)
        chart.add_data(sm_ref, titles_from_data=True)
        chart.add_data(mem_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        if len(chart.series) >= 2:
            chart.series[0].graphicalProperties.solidFill = "C084FC"
            chart.series[1].graphicalProperties.solidFill = "60A5FA"
        style_chart_dark(chart, "SM vs Memory Utilization by Kernel")
        ws3.add_chart(chart, f"A{row}")
        ws3.freeze_panes = "A4"

    # ==================================================================
    #  Sheet 4: Raw Launches
    # ==================================================================
    if kernels:
        ws4 = wb.create_sheet("Raw Launches")
        ws4.sheet_properties.tabColor = "4ADE80"
        prefill_dark(ws4)

        row = 1
        banner(ws4, row, "INDIVIDUAL KERNEL LAUNCHES", 11)
        row += 1
        fill_blank(ws4, row, 11, FILL_DEEP)
        row += 1

        raw_headers = ["#", "Kernel", "Duration (us)", "Grid", "Block",
                       "SM %", "Mem %", "Occ %", "L1 Hit %", "L2 Hit %",
                       "Bottleneck"]
        raw_widths = [5, 34, 14, 14, 12, 9, 9, 9, 10, 10, 28]
        header_row(ws4, row, raw_headers, raw_widths)
        ws4.freeze_panes = f"A{row + 1}"
        row += 1

        sorted_kernels = sorted(kernels, key=lambda k: -get_duration_us(k))
        for i, k in enumerate(sorted_kernels):
            dur = get_duration_us(k)
            bg = FILL_ALT if i % 2 else FILL_BASE
            bn = classify_bottleneck(
                safe_float(k.get("Compute (SM) Throughput")),
                safe_float(k.get("Memory Throughput")),
                safe_float(k.get("Achieved Occupancy")),
            )
            dc(ws4, row, 1, k["id"], fill=bg, font=FONT_DIM, align=ALIGN_C)
            dc(ws4, row, 2, k["name"], fill=bg)
            dc_num(ws4, row, 3, round(dur, 2), fill=bg)
            dc(ws4, row, 4, k.get("grid_size", ""), fill=bg, font=FONT_DIM, align=ALIGN_C)
            dc(ws4, row, 5, k.get("block_size", ""), fill=bg, font=FONT_DIM, align=ALIGN_C)
            for ci, metric in [(6, "Compute (SM) Throughput"),
                               (7, "Memory Throughput"),
                               (8, "Achieved Occupancy"),
                               (9, "L1/TEX Hit Rate"),
                               (10, "L2 Hit Rate")]:
                val = safe_float(k.get(metric))
                display = min(val, 100.0) if ci in (6, 7) else val
                gf, gfn = grade_style(display)
                dc_status(ws4, row, ci, display, (gf, gfn))
            bn_fill, bn_font = bottleneck_style(bn)
            dc(ws4, row, 11, bn, fill=bn_fill, font=bn_font)
            row += 1

        for r in range(row, row + 5):
            fill_blank(ws4, r, 11)

    wb.save(output_path)
    return output_path


# ======================================================================
#  MAIN
# ======================================================================

def main():
    parser = argparse.ArgumentParser(description="FSR4 Executive Summarizer")
    parser.add_argument("ncu_csv", nargs="?", default=None,
                        help="NCU details CSV (from ncu --csv --page details)")
    parser.add_argument("--stdout", help="Benchmark stdout capture file")
    parser.add_argument("-o", "--output", default=None, help="Output .xlsx path")
    args = parser.parse_args()

    kernels = []
    if args.ncu_csv:
        print(f"Parsing NCU data from {args.ncu_csv}...")
        kernels = parse_ncu_csv(args.ncu_csv)
        print(f"  Found {len(kernels)} kernel launches")

    stdout_data = {}
    if args.stdout:
        print(f"Parsing benchmark output from {args.stdout}...")
        stdout_data = parse_fsr4_stdout(args.stdout)
        if stdout_data.get("phases"):
            print(f"  Found {len(stdout_data['phases'])} layer timings")

    if not kernels and not stdout_data:
        print("ERROR: Provide at least one of: ncu_csv or --stdout")
        sys.exit(1)

    output = args.output
    if not output:
        if args.ncu_csv:
            output = str(Path(args.ncu_csv).parent / "fsr4_performance_report.xlsx")
        else:
            output = str(Path(args.stdout).parent / "fsr4_performance_report.xlsx")

    print(f"Building Excel report...")
    build_excel(kernels, stdout_data, output)
    print(f"Report saved to: {output}")


if __name__ == "__main__":
    main()
